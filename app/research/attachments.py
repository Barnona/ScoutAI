"""Utilities for extracting research context from user attachments."""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 24000
ALLOWED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".csv", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp"
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


@dataclass
class AttachmentContext:
    filename: str
    kind: str
    text: str
    size: int


def _trim(text: str) -> str:
    text = text.replace("\x00", " ").strip()
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + "\n[Attachment text truncated.]"
    return text


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:80]:
        parts.append(page.extract_text() or "")
    return _trim("\n\n".join(parts))


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _trim("\n".join(parts))


def _extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[SHEET: {ws.title}]")
        for row in ws.iter_rows(max_row=500, values_only=True):
            values = ["" if v is None else str(v) for v in row]
            if any(values):
                parts.append(" | ".join(values))
    return _trim("\n".join(parts))


def _extract_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for index, slide in enumerate(prs.slides[:100], 1):
        parts.append(f"[SLIDE {index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return _trim("\n".join(parts))


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))[:1000]
    return _trim("\n".join(" | ".join(row) for row in rows))


def extract_attachment(filename: str, data: bytes) -> AttachmentContext:
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported attachment type: {ext or 'unknown'}")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"{filename} exceeds the 10 MB attachment limit.")

    if ext == ".pdf":
        text, kind = _extract_pdf(data), "PDF"
    elif ext == ".docx":
        text, kind = _extract_docx(data), "DOCX"
    elif ext in {".xlsx", ".xls"}:
        text, kind = _extract_xlsx(data), "SPREADSHEET"
    elif ext in {".pptx", ".ppt"}:
        text, kind = _extract_pptx(data), "PRESENTATION"
    elif ext == ".csv":
        text, kind = _extract_csv(data), "CSV"
    elif ext in {".txt", ".md", ".doc"}:
        text, kind = _trim(data.decode("utf-8-sig", errors="replace")), "TEXT"
    elif ext in IMAGE_EXTENSIONS:
        text, kind = "", "IMAGE"
    else:
        text, kind = "", "FILE"

    return AttachmentContext(filename=filename, kind=kind, text=text, size=len(data))


def build_context(attachments: list[AttachmentContext]) -> str:
    parts = []
    for item in attachments:
        if item.text:
            parts.append(
                f"ATTACHED FILE: {item.filename} ({item.kind})\n"
                f"{item.text}"
            )
        elif item.kind == "IMAGE":
            parts.append(
                f"ATTACHED IMAGE: {item.filename}\n"
                "[Image content is available to the multimodal model when supported.]"
            )
    return "\n\n---\n\n".join(parts)
